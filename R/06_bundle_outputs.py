"""06: Bundle everything into TWO files in output/ (needs only Python 3 + openpyxl).

  output/WAR_streaks_report.html   one self-contained page: answers, all 11 charts, the key
                                   tables, and the interactive Streak Explorer (embedded)
  output/WAR_streaks_tables.xlsx   every results table as a sheet, with a clickable contents tab

The individual CSVs / PNGs / explorer page stay in output/parts/ for anyone who wants them.
Run from anywhere:  python R/06_bundle_outputs.py
"""
import base64, csv, html, re, sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
PARTS = ROOT / "output" / "parts"
TABLES, CHARTS = PARTS / "tables", PARTS / "charts"
OUT_HTML = ROOT / "output" / "WAR_streaks_report.html"
OUT_XLSX = ROOT / "output" / "WAR_streaks_tables.xlsx"


def read_csv(p):
    with open(p, encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


# ------------------------------------------------------------------------------------------------
# 1. Excel workbook: one sheet per table, contents tab with links
# ------------------------------------------------------------------------------------------------
DESC = {
    "a1_two_year": ("A1", "Best 2-year stretch"),
    "a2_k3_year": ("A2 3yr", "Best 3-year stretch"), "a2_k5_year": ("A2 5yr", "Best 5-year stretch"),
    "a2_k7_year": ("A2 7yr", "Best 7-year stretch (bonus)"), "a2_k10_year": ("A2 10yr", "Best 10-year stretch (bonus)"),
}
SUFFIX = {"_top30": ("top30", "top 30 players, each at their best window"),
          "_top40_all_overlapping_windows": ("top40 windows", "top 40 windows, overlaps allowed"),
          "_splits": ("splits", "top players inside every split (role, position, decade, age, league, franchise, era cutoffs)"),
          "_sensitivity_modes": ("modes", "strict vs service-adjusted vs pace-adjusted")}
ORDER = (["a1_two_year", "a2_k3_year", "a2_k5_year", "a2_k7_year", "a2_k10_year"])


def sheet_meta(stem):
    for base, (b, bdesc) in DESC.items():
        for suf, (s, sdesc) in SUFFIX.items():
            if stem == base + suf:
                return f"{b} {s}", f"{bdesc}: {sdesc}"
    m = re.match(r"a3_streak_ge(\d+)_top25", stem)
    if m:
        return f"A3 {m.group(1)}+ WAR", f"Longest streaks of seasons at {m.group(1)}+ WAR (top 25 players)"
    fixed = {
        "a3_streak_splits": ("A3 splits", "Longest threshold streaks inside every split"),
        "a3_streak_player_counts": ("A3 counts", "How many players reached each streak length at each threshold"),
        "a3_streak_top10_service_adjusted": ("A3 service-adj", "Threshold streaks with military service bridged"),
        "a4_elite_pairs_top25": ("A4 top25", "Both-seasons-elite pairs, six definitions"),
        "a4_elite_pairs_splits": ("A4 splits", "Both-seasons-elite pairs inside every split"),
        "a4_elite_pairs_counts": ("A4 counts", "How many pairs/players meet each definition"),
        "a4_elite_pairs_sensitivity": ("A4 sensitivity", "Elite pairs: strict vs service-adjusted"),
        "era_cutoff_sensitivity": ("Era cutoffs", "Top 5 either side of the 1901/1920/1947/1969 cutoffs"),
        "peak_vs_longevity": ("Peak vs longevity", "Career WAR vs best 5-year stretch, players with 30+ WAR"),
        "crown_best_window_by_length": ("Crown", "Record WAR for every window length 1-20"),
        "sanity_checks_summary": ("Sanity checks", "All automated checks, pass/fail"),
        "sanity_career_totals": ("Sanity career", "Career WAR vs known Baseball-Reference totals"),
        "sanity_icon_ranks": ("Sanity icons", "Where Ruth, Bonds, Mays... rank for 2/3/5 years"),
        "sanity_traded_players": ("Sanity trades", "Mid-season trades merged into one season"),
        "sanity_two_way_seasons": ("Sanity two-way", "Ruth and Ohtani seasons, hitting vs pitching WAR"),
        "sanity_shortened_seasons_top3": ("Sanity short seasons", "Top 3 in 1981/94/95/2020 raw vs pace-adjusted"),
        "sanity_military_service_effect": ("Sanity service", "Best windows with and without bridging military service"),
        "sanity_best_gap_pairs_not_counted": ("Sanity gap pairs", "Elite pairs NOT counted because a year is missing"),
        "diag_load_summary": ("Diag load", "Row counts from the data load"),
        "diag_season_lengths": ("Diag season length", "Games per team and pace factor by season"),
        "diag_unmatched_war_players": ("Diag unmatched", "WAR players with no Lahman match"),
    }
    return fixed.get(stem, (stem[:31], stem))


def sort_key(p):
    s = p.stem
    sfx = ["_top30", "_top40_all_overlapping_windows", "_splits", "_sensitivity_modes"]
    for i, b in enumerate(ORDER):
        if s.startswith(b):
            return (0, i, [k for k, x in enumerate(sfx) if s.endswith(x)][0], s)
    if s.startswith("a3_"): return (1, 0, 0, s)
    if s.startswith("a4_"): return (2, 0, 0, s)
    for i, pre in enumerate(["era_", "peak_", "crown_", "sanity_", "diag_"]):
        if s.startswith(pre): return (3, i, 0, s)
    return (9, 0, 0, s)


def num(v):
    if v is None or v == "": return None
    try:
        f = float(v)
        return int(f) if re.fullmatch(r"-?\d+", v.strip()) else f
    except ValueError:
        return v


def build_xlsx(files):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.hyperlink import Hyperlink
    wb = Workbook(); idx = wb.active; idx.title = "Contents"
    idx.append(["Sheet (click to jump)", "What it is", "Rows"]); names = []
    head_fill = PatternFill("solid", fgColor="1F3A5F")
    for p in files:
        short, desc = sheet_meta(p.stem)
        name = short[:31]; n = 2
        while name in names: name = f"{short[:28]}_{n}"; n += 1
        names.append(name)
        rows = list(csv.reader(open(p, encoding="utf-8", newline="")))
        ws = wb.create_sheet(name)
        ws.append(rows[0])
        for r in rows[1:]: ws.append([num(c) for c in r])
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = head_fill; c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
        for i, col in enumerate(ws.columns, 1):
            w = max((len(str(c.value)) for c in list(col)[:200] if c.value is not None), default=8)
            ws.column_dimensions[get_column_letter(i)].width = min(max(w + 2, 8), 46)
        idx.append([name, desc, len(rows) - 1])
        link = idx.cell(row=idx.max_row, column=1)
        link.hyperlink = Hyperlink(ref=link.coordinate, location=f"'{name}'!A1", display=name); link.font = Font(color="0563C1", underline="single")
    for c in idx[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = head_fill
    idx.column_dimensions["A"].width = 26; idx.column_dimensions["B"].width = 100; idx.column_dimensions["C"].width = 8
    idx.freeze_panes = "A2"
    wb.save(OUT_XLSX)
    return len(files)


# ------------------------------------------------------------------------------------------------
# 2. HTML report
# ------------------------------------------------------------------------------------------------
def inline(t):
    t = html.escape(t, quote=False)
    t = re.sub(r"`([^`]+)`", r"<code>\1</code>", t)
    t = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", t)
    t = re.sub(r"(?<![\w*])\*([^*\n]+?)\*(?![\w*])", r"<i>\1</i>", t)
    return t


def md_to_html(lines):
    out, i, stack = [], 0, []   # stack of (indent, tag)

    def close_to(ind):
        while stack and stack[-1][0] >= ind:
            out.append(f"</li></{stack.pop()[1]}>")
    while i < len(lines):
        ln = lines[i]
        if ln.strip().startswith("```"):
            close_to(0); i += 1; buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i]); i += 1
            out.append("<pre>" + html.escape("\n".join(buf)) + "</pre>"); i += 1; continue
        if ln.startswith("|"):
            close_to(0); rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells): rows.append(cells)
                i += 1
            out.append("<table class='md'><thead><tr>" + "".join(f"<th>{inline(c)}</th>" for c in rows[0]) + "</tr></thead><tbody>" +
                       "".join("<tr>" + "".join(f"<td>{inline(c)}</td>" for c in r) + "</tr>" for r in rows[1:]) + "</tbody></table>")
            continue
        m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)", ln)
        if m:
            ind = len(m.group(1)); tag = "ol" if m.group(2)[0].isdigit() else "ul"
            if not stack or ind > stack[-1][0]:
                out.append(f"<{tag}>"); stack.append((ind, tag))
            else:
                while stack and stack[-1][0] > ind: out.append(f"</li></{stack.pop()[1]}>")
                out.append("</li>")
            out.append("<li>" + inline(m.group(3))); i += 1; continue
        if ln.strip() == "":
            i += 1; continue
        if stack and ln.startswith(" "):      # continuation of a list item
            out.append(" " + inline(ln.strip())); i += 1; continue
        close_to(0)
        out.append("<p>" + inline(ln.strip()) + "</p>"); i += 1
    close_to(0)
    return "\n".join(out)


def split_sections(md):
    """-> (intro_lines, {h2 title: lines})"""
    secs, cur, intro = {}, None, []
    for ln in md.splitlines():
        if ln.startswith("## "):
            cur = ln[3:].strip(); secs[cur] = []
        elif cur is None:
            if not ln.startswith("# "): intro.append(ln)
        else:
            secs[cur].append(ln)
    return intro, secs


def df_table(rows, cols, labels=None, limit=None):
    labels = labels or cols
    def fmt(v):
        try:
            f = float(v)
            return str(int(f)) if re.fullmatch(r"-?\d+", v.strip()) else f"{f:.1f}"
        except (ValueError, AttributeError):
            return html.escape(v or "")
    rows = rows[:limit] if limit else rows
    th = "".join(f"<th>{html.escape(l)}</th>" for l in labels)
    body = "".join("<tr>" + "".join(f"<td>{fmt(r.get(c, ''))}</td>" for c in cols) + "</tr>" for r in rows)
    return f"<div class='tw'><table class='data'><thead><tr>{th}</tr></thead><tbody>{body}</tbody></table></div>"


CHART_TITLES = {
    "01": "Best two-year stretches", "02": "Best three- and five-year stretches", "03": "Longest streaks above 5, 6 and 8 WAR",
    "04": "When both seasons are elite", "05": "Does the era cutoff change who's best?", "06": "Best five-year stretch by position",
    "07": "Best three-year stretch by decade", "08": "Peak vs. longevity", "09": "Age at the start of elite stretches",
    "10": "Franchises with elite three-year stretches", "11": "The crown: record WAR by window length"}


def build_html():
    md = (ROOT / "README.md").read_text(encoding="utf-8")
    intro, secs = split_sections(md)
    sec = lambda t: md_to_html(secs[t]) if t in secs else ""
    T = lambda n: read_csv(TABLES / f"{n}.csv")
    std = ["rank", "name", "start", "end", "WAR_sum", "war_seq", "role_label", "pos_group", "franch_peak"]
    stl = ["#", "Player", "From", "To", "WAR", "Season-by-season WAR", "Role", "Pos", "Team"]
    details = []

    def det(title, body, open_=False):
        details.append(f"<details{' open' if open_ else ''}><summary>{html.escape(title)}</summary>{body}</details>")
    det("Best 2-year stretches (top 30 players)", df_table(T("a1_two_year_top30"), std, stl), True)
    det("Best 3-year stretches (top 30)", df_table(T("a2_k3_year_top30"), std, stl))
    det("Best 5-year stretches (top 30)", df_table(T("a2_k5_year_top30"), std, stl))
    det("Best 7-year stretches (top 30, bonus)", df_table(T("a2_k7_year_top30"), std, stl))
    det("Best 10-year stretches (top 30, bonus)", df_table(T("a2_k10_year_top30"), std, stl))
    for thr in (5, 6, 8):
        det(f"Longest streaks of {thr}+ WAR seasons (top 25)",
            df_table(T(f"a3_streak_ge{thr}_top25"), ["rank", "name", "start", "end", "k", "WAR_sum", "role_label", "pos_group"],
                     ["#", "Player", "From", "To", "Seasons", "Total WAR", "Role", "Pos"]))
    a4 = T("a4_elite_pairs_top25"); defs = []
    for r in a4:
        if r["definition"] not in defs: defs.append(r["definition"])
    for d in defs:
        det(f"Both seasons elite: {d} (top 15)", df_table([r for r in a4 if r["definition"] == d], std, stl, 15))
    det("Era cutoff sensitivity (top 5 either side of each cutoff, 2- and 5-year)",
        df_table(T("era_cutoff_sensitivity"), ["k", "cutoff", "side", "rank", "name", "start", "end", "WAR_sum"],
                 ["Years", "Cutoff", "Side", "#", "Player", "From", "To", "WAR"]))
    pl = T("peak_vs_longevity")[:80]
    det("Peak vs. longevity (top 80 careers)", df_table(pl, ["name", "first_year", "last_year", "career_WAR", "best5", "peak5_share", "career_type"],
                                                       ["Player", "Debut", "Last", "Career WAR", "Best 5-yr", "Share of career in best 5", "Type"]))
    det("The crown: record holders for every window length",
        df_table([r for r in T("crown_best_window_by_length") if r["rank"] == "1"], ["k", "name", "start", "end", "WAR_sum"],
                 ["Seasons", "Record holder", "From", "To", "WAR"]))
    chk = T("sanity_checks_summary")
    det("Sanity checks (all automated)", df_table(chk, ["check", "result", "pass"], ["Check", "Result", "Passed"]))

    figs = []
    for p in sorted(CHARTS.glob("*.png")):
        n = p.name[:2]
        b64 = base64.b64encode(p.read_bytes()).decode()
        figs.append(f"<figure><figcaption>{html.escape(CHART_TITLES.get(n, p.stem))}</figcaption>"
                    f"<img alt='{html.escape(CHART_TITLES.get(n, p.stem))}' src='data:image/png;base64,{b64}'></figure>")

    expl = (PARTS / "streak_explorer.html").read_text(encoding="utf-8")
    srcdoc = expl.replace("&", "&amp;").replace('"', "&quot;")
    n_tables = len(list(TABLES.glob("*.csv")))

    css = """
:root{--bg:#fafaf7;--card:#fff;--ink:#1a1d23;--muted:#667085;--line:#e3e5ea;--acc:#0f766e;--th:#1f3a5f}
@media (prefers-color-scheme:dark){:root{--bg:#0f1218;--card:#171b24;--ink:#e8eaf0;--muted:#9aa1b2;--line:#2a303d;--acc:#2dd4bf;--th:#27406a}}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--ink);font:16px/1.55 system-ui,-apple-system,Segoe UI,Roboto,sans-serif}
header{padding:28px 20px 10px;max-width:1100px;margin:0 auto}h1{margin:0 0 6px;font-size:30px;letter-spacing:-.4px}
nav{position:sticky;top:0;background:color-mix(in srgb,var(--bg) 92%,transparent);backdrop-filter:blur(6px);border-bottom:1px solid var(--line);z-index:5}
nav div{max-width:1100px;margin:0 auto;padding:8px 20px;display:flex;gap:16px;flex-wrap:wrap;font-size:14px}
nav a{color:var(--acc);text-decoration:none;font-weight:600}
main{max-width:1100px;margin:0 auto;padding:0 20px 80px}
section{margin-top:34px;scroll-margin-top:50px}h2{font-size:23px;border-bottom:2px solid var(--line);padding-bottom:6px;margin-bottom:12px}
p,li{max-width:80ch}ul,ol{padding-left:22px}li{margin:3px 0}code{background:color-mix(in srgb,var(--ink) 9%,transparent);padding:1px 5px;border-radius:4px;font-size:.9em}
table{border-collapse:collapse;font-size:14px;margin:10px 0}th{background:var(--th);color:#fff;text-align:left;padding:6px 10px;font-weight:600;white-space:nowrap}
td{padding:5px 10px;border-top:1px solid var(--line);vertical-align:top}.tw{overflow-x:auto}table.data td:nth-child(n+3){white-space:nowrap}
figure{margin:18px 0;background:var(--card);border:1px solid var(--line);border-radius:12px;padding:12px}figcaption{font-weight:700;margin-bottom:8px}
img{width:100%;height:auto;border-radius:6px;background:#fff}
details{background:var(--card);border:1px solid var(--line);border-radius:10px;margin:8px 0;padding:0 14px}summary{cursor:pointer;padding:10px 0;font-weight:600}
iframe{width:100%;height:1400px;border:1px solid var(--line);border-radius:12px;background:var(--bg)}
.note{color:var(--muted);font-size:14px}pre{background:var(--card);border:1px solid var(--line);padding:10px;border-radius:8px;overflow:auto}
"""
    resize = """<script>
(function(){var f=document.getElementById('xf');function fit(){try{var d=f.contentDocument;if(d&&d.body)f.style.height=(d.body.scrollHeight+24)+'px';}catch(e){}}
f.addEventListener('load',function(){fit();setInterval(fit,700);});})();
</script>"""
    page = f"""<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>MLB WAR Streaks Report</title><style>{css}</style></head><body>
<header><h1>Who has the highest WAR in consecutive seasons?</h1>
<div class="note">MLB history, 1876&ndash;2024 &middot; Baseball-Reference WAR joined to Lahman 13.0.0 &middot; one-file report (charts, tables and the interactive explorer are all inside this page)</div></header>
<nav><div><a href="#answers">Answers</a><a href="#splits">Splits</a><a href="#surprises">Surprises</a><a href="#charts">Charts</a><a href="#explorer">Explorer</a><a href="#tables">Tables</a><a href="#method">Method &amp; caveats</a></div></nav>
<main>
<section id="answers"><h2>Headline answers</h2>{md_to_html(intro)}{sec("Headline answers")}</section>
<section id="splits"><h2>TL;DR for each split</h2>{sec("TL;DR for each split")}</section>
<section id="surprises"><h2>Anything surprising</h2>{sec("Anything surprising")}</section>
<section id="charts"><h2>Charts</h2>{''.join(figs)}</section>
<section id="explorer"><h2>WAR Streak Explorer (interactive)</h2>
<p class="note">The extra feature I added. Search a player, drag the window length, see their best run, rank, and streak twins.</p>
<iframe id="xf" title="WAR Streak Explorer" srcdoc="{srcdoc}"></iframe>{sec("The unexpected extra: WAR Streak Explorer")}</section>
<section id="tables"><h2>Key tables</h2>
<p class="note">Click a heading to expand. The complete set ({n_tables} tables, including every split) is in <b>WAR_streaks_tables.xlsx</b>, next to this file.</p>{''.join(details)}</section>
<section id="method"><h2>Method &amp; caveats</h2>
<h3>Where the WAR comes from</h3>{sec("Where the WAR comes from (and the caveats)")}
<h3>Definitions</h3>{sec("Definitions")}
<h3>Sanity checks</h3>{sec("Sanity checks")}</section>
</main>{resize}</body></html>"""
    OUT_HTML.write_text(page, encoding="utf-8")
    return OUT_HTML.stat().st_size / 1e6


if __name__ == "__main__":
    if not TABLES.exists():
        sys.exit(f"Run the R pipeline first (missing {TABLES})")
    csvs = sorted(TABLES.glob("*.csv"), key=sort_key)
    n = build_xlsx(csvs)
    print(f"Wrote {OUT_XLSX.name}: {n} sheets")
    print(f"Wrote {OUT_HTML.name}: {build_html():.1f} MB")
